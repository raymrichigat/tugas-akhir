"""
Convert bimbingan.md -> bimbingan.docx
Custom markdown parser yang handle:
  - Headings (#, ##, ###)
  - Tables (| ... | ... |)
  - Code blocks (```)
  - Blockquotes (>)
  - Numbered & bullet lists
  - Inline: **bold**, *italic*, `code`, [text](url)
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).parent
MD_PATH = ROOT / "bimbingan.md"
DOCX_PATH = ROOT / "bimbingan.docx"


# ─── Inline formatting ──────────────────────────────────────────────────────────

INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*)"          # **bold**
    r"|(\*[^*]+\*)"             # *italic*
    r"|(`[^`]+`)"               # `code`
    r"|(\[[^\]]+\]\([^)]+\))"   # [text](url)
)


def add_inline_runs(paragraph, text: str):
    """Tambahkan runs ke paragraph dengan inline formatting."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])

        bold_m, italic_m, code_m, link_m = m.groups()
        if bold_m:
            run = paragraph.add_run(bold_m[2:-2])
            run.bold = True
        elif italic_m:
            run = paragraph.add_run(italic_m[1:-1])
            run.italic = True
        elif code_m:
            run = paragraph.add_run(code_m[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif link_m:
            mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", link_m)
            label, url = mm.group(1), mm.group(2)
            run = paragraph.add_run(label)
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            run.underline = True
            # Inject hyperlink (simple approach via XML)
            try:
                add_hyperlink(paragraph, run, url)
            except Exception:
                pass
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_hyperlink(paragraph, run, url):
    """Wrap a run di dalam <w:hyperlink>."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    r_element = run._r
    r_element.getparent().remove(r_element)
    hyperlink.append(r_element)
    paragraph._p.append(hyperlink)


# ─── Block parser ──────────────────────────────────────────────────────────────

def make_doc():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    # Default body font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    return doc


def parse_table_block(lines):
    """lines: list of '| a | b |' strings (sudah dibuang separator '---')."""
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def add_table_to_doc(doc, header_row, body_rows):
    n_cols = len(header_row)
    if n_cols == 0:
        return

    table = doc.add_table(rows=1 + len(body_rows), cols=n_cols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Header
    for i, cell_text in enumerate(header_row):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        add_inline_runs(p, cell_text)
        for run in p.runs:
            run.bold = True

    # Body
    for r, row in enumerate(body_rows, start=1):
        for c, cell_text in enumerate(row):
            if c >= n_cols:
                break
            cell = table.rows[r].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline_runs(p, cell_text)


def add_code_block(doc, code_text, lang=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)

    # Add a light gray shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F4F4F4")
    pPr.append(shd)


def add_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    if level == 1:
        run.font.size = Pt(20)
    elif level == 2:
        run.font.size = Pt(15)
    else:
        run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


# ─── Main ──────────────────────────────────────────────────────────────────────

def md_to_docx(md_text: str, doc):
    lines = md_text.splitlines()
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.rstrip()

        # ─── Code block ────────────────────────
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            i += 1
            code_lines = []
            while i < n and not lines[i].rstrip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            add_code_block(doc, "\n".join(code_lines), lang)
            continue

        # ─── Heading ───────────────────────────
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            add_heading(doc, text, level=min(level, 4))
            i += 1
            continue

        # ─── Horizontal rule ───────────────────
        if re.match(r"^-{3,}$|^\*{3,}$|^_{3,}$", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("─" * 60)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            i += 1
            continue

        # ─── Table ─────────────────────────────
        if stripped.startswith("|") and i + 1 < n:
            sep = lines[i + 1].strip()
            if re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$", sep):
                # Header
                header_line = stripped
                table_lines = []
                i += 2  # skip header + separator
                while i < n and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i])
                    i += 1
                header_row = parse_table_block([header_line])[0]
                body_rows = parse_table_block(table_lines)
                add_table_to_doc(doc, header_row, body_rows)
                doc.add_paragraph()  # spacing setelah tabel
                continue

        # ─── Blockquote ────────────────────────
        if stripped.startswith("> "):
            quote_lines = []
            while i < n and lines[i].rstrip().startswith(">"):
                quote_lines.append(lines[i].lstrip("> ").rstrip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            text = " ".join(quote_lines)
            add_inline_runs(p, text)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue

        # ─── Numbered list ─────────────────────
        m_num = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m_num:
            text = m_num.group(3)
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, text)
            i += 1
            continue

        # ─── Bullet list ───────────────────────
        m_bul = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m_bul:
            text = m_bul.group(2)
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, text)
            i += 1
            continue

        # ─── Empty line ────────────────────────
        if not stripped:
            i += 1
            continue

        # ─── Paragraph (default) ───────────────
        # Gabung baris berturut sampai ketemu blank/heading/list/table
        para_lines = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].rstrip()
            if (not nxt
                or nxt.startswith("#")
                or nxt.startswith("```")
                or nxt.startswith(">")
                or re.match(r"^(\s*)(\d+)\.\s+", lines[i])
                or re.match(r"^(\s*)[-*+]\s+", lines[i])
                or nxt.startswith("|")):
                break
            para_lines.append(nxt)
            i += 1
        text = " ".join(para_lines)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline_runs(p, text)


def main():
    md_text = MD_PATH.read_text(encoding="utf-8")
    doc = make_doc()
    md_to_docx(md_text, doc)
    doc.save(DOCX_PATH)
    print(f"OK -> {DOCX_PATH}")
    print(f"Size: {DOCX_PATH.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
